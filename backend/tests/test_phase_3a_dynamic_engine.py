from io import BytesIO
import pytest
from pptx import Presentation
from app.schemas.recognition import RecognitionRecord
from app.schemas.mapping import FieldMappingConfig, FieldMappingDetail, SlotMappingDetail
from app.services.pptx_generator import generate_powerpoint_presentation
from app.services.template_inspector import inspect_powerpoint_template


def create_multi_card_pptx(slots_per_slide: int = 4) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Add text boxes and picture placeholders for N cards
    for i in range(1, slots_per_slide + 1):
        left_offset = (i - 1) * 2.5
        name_box = slide.shapes.add_textbox(left_offset, 1.0, 2.0, 0.5)
        name_box.name = f"Employee_Name_{i}"
        name_box.text_frame.text = f"Sample Name {i}"
        
        desig_box = slide.shapes.add_textbox(left_offset, 1.6, 2.0, 0.5)
        desig_box.name = f"Designation_{i}"
        desig_box.text_frame.text = f"Sample Role {i}"

        photo_box = slide.shapes.add_textbox(left_offset, 2.2, 2.0, 1.5)
        photo_box.name = f"Photo_Placeholder_{i}"
        photo_box.text_frame.text = "[Insert Photo Here]"

    bio = BytesIO()
    prs.save(bio)
    return bio.getvalue()


def test_multi_card_layout_inspection():
    pptx_bytes = create_multi_card_pptx(4)
    inspection = inspect_powerpoint_template(pptx_bytes, "multi_card.pptx")
    assert inspection.valid is True
    assert inspection.slide_count == 1
    layout = inspection.slides[0].layout_model
    assert layout is not None
    assert layout.capacity == 4
    assert len(layout.employee_slots) == 4


def test_multi_card_generation_overflow():
    pptx_bytes = create_multi_card_pptx(4)
    records = [
        RecognitionRecord(
            row_number=i,
            employee_name=f"User {i}",
            designation=f"Engineer {i}",
            branch="Helsinki",
            award_name="Excellence",
            is_valid=True,
        )
        for i in range(1, 10)  # 9 records across 4-card capacity -> 3 slides
    ]
    
    slots = [
        SlotMappingDetail(
            slot_index=i-1,
            employee_name=FieldMappingDetail(shape_name=f"Employee_Name_{i}"),
            designation=FieldMappingDetail(shape_name=f"Designation_{i}"),
            photo_placeholder=FieldMappingDetail(shape_name=f"Photo_Placeholder_{i}"),
        )
        for i in range(1, 5)
    ]
    mapping = FieldMappingConfig(template_slide_index=0, slots=slots)

    res = generate_powerpoint_presentation(
        template_content=pptx_bytes,
        template_filename="multi_card.pptx",
        excel_records=records,
        mapping_config=mapping,
        source_excel_file_id="test_excel",
        source_template_file_id="test_template",
    )

    assert res.record_count == 9
    assert res.slide_count == 3  # 4 + 4 + 1 employee record
    assert res.status == "completed"
