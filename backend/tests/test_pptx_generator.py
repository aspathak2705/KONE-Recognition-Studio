import pytest
from pptx import Presentation
from io import BytesIO
from app.schemas.recognition import RecognitionRecord
from app.schemas.mapping import FieldMappingConfig, FieldMappingDetail
from app.services.pptx_generator import generate_powerpoint_presentation
from app.services.template_inspector import inspect_powerpoint_template


def create_mock_template_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    
    # Add text shapes for dynamic fields
    emp_shape = slide.shapes.add_textbox(100, 100, 300, 50)
    emp_shape.name = "EmployeeName"
    emp_shape.text_frame.text = "Sample Name"

    des_shape = slide.shapes.add_textbox(100, 160, 300, 50)
    des_shape.name = "Designation"
    des_shape.text_frame.text = "Sample Role"

    br_shape = slide.shapes.add_textbox(100, 220, 300, 50)
    br_shape.name = "Branch"
    br_shape.text_frame.text = "Sample City"

    aw_shape = slide.shapes.add_textbox(100, 280, 300, 50)
    aw_shape.name = "AwardName"
    aw_shape.text_frame.text = "Sample Award"

    bio = BytesIO()
    prs.save(bio)
    return bio.getvalue()


def test_generate_powerpoint_presentation():
    tpl_bytes = create_mock_template_pptx()
    records = [
        RecognitionRecord(
            row_number=2,
            employee_name="Alice Smith",
            designation="Senior Developer",
            branch="Pune",
            award_name="Excellence Award",
        ),
        RecognitionRecord(
            row_number=3,
            employee_name="Bob Jones",
            designation="Lead Architect",
            branch="Mumbai",
            award_name="Star Performer",
        ),
    ]

    cfg = FieldMappingConfig(
        employee_name=FieldMappingDetail(shape_name="EmployeeName"),
        designation=FieldMappingDetail(shape_name="Designation"),
        branch=FieldMappingDetail(shape_name="Branch"),
        award_name=FieldMappingDetail(shape_name="AwardName"),
    )

    res = generate_powerpoint_presentation(
        template_content=tpl_bytes,
        template_filename="test_template.pptx",
        excel_records=records,
        mapping_config=cfg,
        source_excel_file_id="exc123",
        source_template_file_id="tpl123",
    )

    assert res.status == "completed"
    assert res.record_count == 2
    assert res.slide_count == 2
    assert res.generation_id is not None
    assert "/api/generations/" in res.download_url
