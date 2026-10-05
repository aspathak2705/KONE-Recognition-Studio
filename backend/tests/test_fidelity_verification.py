import io
import json
import pytest
from pathlib import Path
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.fidelity import TemplateFidelityManifest, VerificationStatus
from app.services.template_registry_service import template_registry
from app.services.template_fidelity_validator import TemplateFidelityValidator
from app.core.config import settings

client = TestClient(app)


def test_fidelity_validator_detects_defective_pptx():
    """Negative regression test: Verify that defective PPTX missing photo regions is detected and BLOCKED."""
    # Find New Joiner registered template
    templates = template_registry.list_templates()
    nj_tpl = next((t for t in templates if "New Joiner" in t.name), None)
    assert nj_tpl is not None, "New Joiner template must be registered"

    manifest = template_registry.get_template_fidelity_manifest(nj_tpl.template_id, nj_tpl.current_version)
    assert manifest is not None, "Fidelity manifest must exist for registered New Joiner template"

    # Defective generated file path
    bad_pptx_path = Path(r"C:\Users\athar\Downloads\KONE_Recognition_e3e01648.pptx")
    if not bad_pptx_path.exists():
        pytest.skip("Reference bad PPTX fixture not found in Downloads")

    with open(bad_pptx_path, "rb") as f:
        bad_bytes = f.read()

    # Run validation
    report = TemplateFidelityValidator.validate_presentation(
        generated_pptx_bytes=bad_bytes,
        manifest=manifest,
        expected_records_count=34,
    )

    # Must be BLOCKED due to missing photo regions
    assert report.verification_status == VerificationStatus.BLOCKED
    assert report.photo_slot_result == "FAIL"
    photo_failures = [f for f in report.failures if f.failure_type == "MISSING_PHOTO_REGION"]
    assert len(photo_failures) > 0, "Validator must flag MISSING_PHOTO_REGION on defective generated output"


def test_download_gate_blocks_unverified_generation():
    """Verify that GET /api/generations/{id}/download rejects downloads when validation_status is BLOCKED."""
    fake_gen_id = "a1b2c3d4e5f6a1b2c3d4e5f6a1b2c3d4"
    settings.GENERATED_OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    
    # Write mock pptx
    fake_pptx = settings.GENERATED_OUTPUTS_DIR / f"{fake_gen_id}.pptx"
    fake_pptx.write_bytes(b"PK\x03\x04mockpptx")

    # Write mock metadata with BLOCKED status
    fake_meta = settings.GENERATED_OUTPUTS_DIR / f"{fake_gen_id}.json"
    with open(fake_meta, "w", encoding="utf-8") as f:
        json.dump({
            "generation_id": fake_gen_id,
            "status": "completed",
            "validation_status": "BLOCKED",
            "validation_score": 0.4,
        }, f)

    res = client.get(f"/api/generations/{fake_gen_id}/download")
    assert res.status_code == 403
    data = res.json()
    assert "detail" in data
    assert data["detail"]["error"] == "PRESENTATION_NOT_VERIFIED"

    # Cleanup mock files
    fake_pptx.unlink(missing_ok=True)
    fake_meta.unlink(missing_ok=True)


def test_validation_endpoint_returns_report():
    """Verify that GET /api/generations/{id}/validation returns the persisted report."""
    fake_gen_id = "b2c3d4e5f6a1b2c3d4e5f6a1b2c3d4e5"
    settings.GENERATED_OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

    fake_report_path = settings.GENERATED_OUTPUTS_DIR / f"{fake_gen_id}_validation.json"
    with open(fake_report_path, "w", encoding="utf-8") as f:
        json.dump({
            "generation_id": fake_gen_id,
            "verification_status": "VERIFIED",
            "overall_score": 1.0,
            "slides_checked": 2,
            "slides_passed": 2,
            "slides_failed": 0,
            "failures": [],
            "warnings": [],
        }, f)

    res = client.get(f"/api/generations/{fake_gen_id}/validation")
    assert res.status_code == 200
    data = res.json()
    assert data["generation_id"] == fake_gen_id
    assert data["verification_status"] == "VERIFIED"

    # Cleanup
    fake_report_path.unlink(missing_ok=True)


def test_generic_unseen_template_pipeline(tmp_path):
    """Verify that a completely generic, unseen PPTX fixture works through register -> manifest -> generate -> validate -> download gate."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from app.schemas.recognition import RecognitionRecord
    from app.schemas.mapping import FieldMappingConfig, FieldMappingDetail, SlotMappingDetail, CardLayoutConfig
    from app.services.pptx_generator import generate_powerpoint_presentation

    # Create unseen PPTX fixture: custom 4:3 slide, unusual layout, custom non-KONE shapes
    custom_prs = Presentation()
    custom_prs.slide_width = Inches(10)
    custom_prs.slide_height = Inches(7.5)
    blank_layout = custom_prs.slide_layouts[6]
    slide = custom_prs.slides.add_slide(blank_layout)

    # Static title banner
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9.0), Inches(1.0))
    tb.text_frame.text = "Acme Global Innovation Honors"

    # Dynamic employee card
    card_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.0), Inches(8.0), Inches(4.0))
    card_box.name = "HonorCard_1"
    tf = card_box.text_frame
    tf.text = "Candidate Name Here\nSenior Specialist\nGlobal Technology"

    custom_bytes_io = io.BytesIO()
    custom_prs.save(custom_bytes_io)
    custom_bytes = custom_bytes_io.getvalue()

    # 1. Register unseen template
    reg_resp = template_registry.register_or_update_template("Acme Honors Template", custom_bytes, "acme_honors.pptx")
    unseen_tid = reg_resp.template.template_id
    unseen_ver = reg_resp.active_version.version_number

    try:
        # 2. Retrieve discovered manifest
        manifest = template_registry.get_template_fidelity_manifest(unseen_tid, unseen_ver)
        assert manifest is not None
        assert manifest.slide_width_inches == 10.0
        assert manifest.slide_height_inches == 7.5
        assert len(manifest.layouts) > 0

        # 3. Create mapping derived from inspection
        mapping = FieldMappingConfig(
            employee_name=FieldMappingDetail(shape_name="HonorCard_1"),
            designation=FieldMappingDetail(shape_name="HonorCard_1"),
            branch=FieldMappingDetail(shape_name="HonorCard_1"),
            slots=[
                SlotMappingDetail(
                    slot_index=0,
                    employee_name=FieldMappingDetail(shape_name="HonorCard_1"),
                    designation=FieldMappingDetail(shape_name="HonorCard_1"),
                    branch=FieldMappingDetail(shape_name="HonorCard_1"),
                )
            ],
            layout_config=CardLayoutConfig(cards_per_slide=1)
        )
        template_registry.update_template_mapping(unseen_tid, mapping, unseen_ver)

        # 4. Generate presentation
        rec = RecognitionRecord(
            row_number=2,
            employee_name="Dr. Jane Smith",
            designation="Principal AI Scientist",
            branch="Tokyo Laboratory",
            award_name="Innovation Champion",
            is_valid=True,
            errors=[]
        )
        gen_res = generate_powerpoint_presentation(
            template_content=custom_bytes,
            template_filename="acme_honors.pptx",
            excel_records=[rec],
            mapping_config=mapping,
            source_excel_file_id="11112222333344445555666677778888",
            source_template_file_id=unseen_tid,
        )

        # Run TemplateFidelityValidator
        gen_path = settings.GENERATED_OUTPUTS_DIR / f"{gen_res.generation_id}.pptx"
        with open(gen_path, "rb") as f:
            generated_bytes = f.read()

        val_report = TemplateFidelityValidator.validate_presentation(
            generated_pptx_bytes=generated_bytes,
            manifest=manifest,
            expected_records_count=1,
        )

        assert gen_res.slide_count >= 1
        assert val_report.verification_status == VerificationStatus.VERIFIED
        assert val_report.overall_score >= 0.8

        # 5. Persist and verify download gate allows VERIFIED output
        meta_path = settings.GENERATED_OUTPUTS_DIR / f"{gen_res.generation_id}.json"
        with open(meta_path, "w", encoding="utf-8") as f:
            import json
            json.dump({
                "generation_id": gen_res.generation_id,
                "status": "completed",
                "validation_status": "VERIFIED",
                "validation_score": val_report.overall_score,
            }, f)

        dl_res = client.get(f"/api/generations/{gen_res.generation_id}/download")
        assert dl_res.status_code == 200
        assert len(dl_res.content) > 0

    finally:
        template_registry.delete_template(unseen_tid)


def test_cross_template_and_version_isolation():
    """Verify that template and version manifests remain isolated and do not cross-contaminate."""
    templates = template_registry.list_templates()
    if len(templates) < 2:
        pytest.skip("Requires at least 2 templates in registry for cross-isolation test")

    t1 = templates[0]
    t2 = templates[1]

    m1 = template_registry.get_template_fidelity_manifest(t1.template_id, t1.current_version)
    m2 = template_registry.get_template_fidelity_manifest(t2.template_id, t2.current_version)

    assert m1 is not None and m2 is not None
    assert m1.template_id == t1.template_id
    assert m2.template_id == t2.template_id
    assert m1.template_id != m2.template_id
    assert m1.manifest_id != m2.manifest_id
