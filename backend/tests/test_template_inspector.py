import pytest
from pptx import Presentation
from io import BytesIO
from app.services.template_inspector import inspect_powerpoint_template


def create_mock_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    title.text = "Sample Presentation Title"
    
    subtitle = slide.placeholders[1]
    subtitle.text = "Sample Subtitle Text"

    bio = BytesIO()
    prs.save(bio)
    return bio.getvalue()


def test_pptx_inspection():
    content = create_mock_pptx()
    res = inspect_powerpoint_template(content, "sample.pptx")

    assert res.valid is True
    assert res.slide_count == 1
    assert res.slide_width_inches > 0
    assert res.slide_height_inches > 0
    assert res.aspect_ratio in ["16:9", "4:3"]
    assert len(res.slides) == 1
    assert res.slides[0].placeholders_count == 2
    assert len(res.slides[0].text_shapes) >= 2
    assert res.slides[0].text_shapes[0].text == "Sample Presentation Title"


def test_corrupted_pptx():
    res = inspect_powerpoint_template(b"not a pptx", "broken.pptx")
    assert res.valid is False
    assert res.slide_count == 0
    assert "Corrupted or invalid" in res.warnings[0]
