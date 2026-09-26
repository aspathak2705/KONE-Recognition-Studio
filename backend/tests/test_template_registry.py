import pytest
from io import BytesIO
from pptx import Presentation
from app.services.template_registry_service import TemplateRegistryService, calculate_pptx_hash
from app.schemas.mapping import FieldMappingConfig, FieldMappingDetail, GenerationReadiness
from app.schemas.template_registry import TemplateMetadata


def create_mock_pptx_bytes(title_text: str = "Employee Name") -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx_box = slide.shapes.add_textbox(0, 0, 100, 50)
    tx_box.name = "EmployeeName"
    tx_box.text_frame.text = title_text
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


@pytest.fixture
def tmp_registry(tmp_path):
    reg_dir = tmp_path / "templates"
    return TemplateRegistryService(registry_dir=reg_dir)


def test_register_new_template_and_hash(tmp_registry):
    content = create_mock_pptx_bytes("John Doe")
    expected_hash = calculate_pptx_hash(content)

    res = tmp_registry.register_or_update_template(
        name="KONE Quarterly Recognition",
        content=content,
        filename="kone_master.pptx",
    )

    assert res.is_new_version is True
    assert res.is_duplicate_hash is False
    assert res.template.name == "KONE Quarterly Recognition"
    assert res.template.current_version == 1
    assert res.template.file_hash == expected_hash
    assert len(res.template.versions) == 1
    assert res.active_version.version_number == 1


def test_duplicate_pptx_hash_reuses_version(tmp_registry):
    content = create_mock_pptx_bytes("Identical Binary")

    res1 = tmp_registry.register_or_update_template(
        name="KONE Master",
        content=content,
        filename="master.pptx",
    )

    res2 = tmp_registry.register_or_update_template(
        name="KONE Master Second Upload",
        content=content,
        filename="master_copy.pptx",
    )

    assert res2.is_duplicate_hash is True
    assert res2.is_new_version is False
    assert res2.template.template_id == res1.template.template_id
    assert res2.active_version.version_number == 1


def test_new_version_created_on_different_hash(tmp_registry):
    content1 = create_mock_pptx_bytes("Version 1 Text")
    content2 = create_mock_pptx_bytes("Version 2 Modified Text")

    res1 = tmp_registry.register_or_update_template(
        name="KONE Award Template",
        content=content1,
        filename="v1.pptx",
    )

    res2 = tmp_registry.register_or_update_template(
        name="KONE Award Template",
        content=content2,
        filename="v2.pptx",
        existing_template_id=res1.template.template_id,
    )

    assert res2.is_new_version is True
    assert res2.template.current_version == 2
    assert len(res2.template.versions) == 2
    assert res2.active_version.version_number == 2
    assert res2.template.versions[0].file_hash != res2.template.versions[1].file_hash


def test_mapping_persistence(tmp_registry):
    content = create_mock_pptx_bytes("Test Mapping")
    res = tmp_registry.register_or_update_template(
        name="Mapping Test Template",
        content=content,
        filename="mapping_test.pptx",
    )

    new_mapping = FieldMappingConfig(
        employee_name=FieldMappingDetail(shape_name="EmployeeName", required=True),
        designation=FieldMappingDetail(shape_name="Designation", required=False),
        branch=FieldMappingDetail(shape_name="Branch", required=False),
        award_name=FieldMappingDetail(shape_name="AwardName", required=False),
    )

    updated_meta = tmp_registry.update_template_mapping(
        template_id=res.template.template_id,
        mapping=new_mapping,
        version_number=1,
    )

    assert updated_meta.versions[0].mapping_config.employee_name.shape_name == "EmployeeName"

    # Verify reloading from disk maintains mapping
    reloaded = tmp_registry.get_template(res.template.template_id)
    assert reloaded.versions[0].mapping_config.employee_name.shape_name == "EmployeeName"


def test_archive_template(tmp_registry):
    content = create_mock_pptx_bytes("Archive Test")
    res = tmp_registry.register_or_update_template(
        name="To Archive",
        content=content,
        filename="archive.pptx",
    )

    archived = tmp_registry.archive_template(res.template.template_id)
    assert archived.is_archived is True

    all_active = tmp_registry.list_templates(include_archived=False)
    assert len(all_active) == 0

    all_with_archived = tmp_registry.list_templates(include_archived=True)
    assert len(all_with_archived) == 1


def test_corrupt_pptx_registration_fails(tmp_registry):
    corrupt_bytes = b"Not a real PPTX presentation binary content"
    with pytest.raises(Exception) as exc_info:
        tmp_registry.register_or_update_template(
            name="Corrupt Template",
            content=corrupt_bytes,
            filename="corrupt.pptx",
        )
    assert "registration failed" in str(exc_info.value.detail).lower()


def test_version_immutability(tmp_registry):
    content1 = create_mock_pptx_bytes("Original v1 text")
    content2 = create_mock_pptx_bytes("Modified v2 text")

    res1 = tmp_registry.register_or_update_template(
        name="Immutable Test",
        content=content1,
        filename="v1.pptx",
    )

    mapping_v1 = FieldMappingConfig(
        employee_name=FieldMappingDetail(shape_name="EmployeeName", required=True)
    )
    tmp_registry.update_template_mapping(res1.template.template_id, mapping_v1, version_number=1)

    res2 = tmp_registry.register_or_update_template(
        name="Immutable Test",
        content=content2,
        filename="v2.pptx",
        existing_template_id=res1.template.template_id,
    )

    mapping_v2 = FieldMappingConfig(
        employee_name=FieldMappingDetail(shape_name="EmployeeName_V2", required=True)
    )
    tmp_registry.update_template_mapping(res1.template.template_id, mapping_v2, version_number=2)

    # Reload from disk
    meta = tmp_registry.get_template(res1.template.template_id)
    assert len(meta.versions) == 2
    assert meta.versions[0].mapping_config.employee_name.shape_name == "EmployeeName"
    assert meta.versions[1].mapping_config.employee_name.shape_name == "EmployeeName_V2"
    assert meta.versions[0].file_hash != meta.versions[1].file_hash


def test_registry_restart_persistence(tmp_path):
    reg_dir = tmp_path / "persistent_templates"
    reg1 = TemplateRegistryService(registry_dir=reg_dir)

    content = create_mock_pptx_bytes("Restart persistence test")
    res = reg1.register_or_update_template(
        name="Persistent Template",
        content=content,
        filename="persistent.pptx",
    )

    # Instantiate new TemplateRegistryService reading from same directory (simulating restart)
    reg2 = TemplateRegistryService(registry_dir=reg_dir)
    reloaded = reg2.get_template(res.template.template_id)

    assert reloaded is not None
    assert reloaded.name == "Persistent Template"
    assert reloaded.file_hash == res.template.file_hash
    assert reloaded.generation_readiness == res.template.generation_readiness


def test_generation_job_version_binding(tmp_registry):
    content1 = create_mock_pptx_bytes("v1 text")
    res1 = tmp_registry.register_or_update_template(
        name="Job Binding Test",
        content=content1,
        filename="v1.pptx",
    )

    # Fetch PPTX for version 1
    _, v1_obj = tmp_registry.get_template_version_pptx(res1.template.template_id, version_number=1)
    assert v1_obj.version_number == 1
    assert v1_obj.file_hash == res1.active_version.file_hash

    # Create version 2
    content2 = create_mock_pptx_bytes("v2 text")
    res2 = tmp_registry.register_or_update_template(
        name="Job Binding Test",
        content=content2,
        filename="v2.pptx",
        existing_template_id=res1.template.template_id,
    )

    # Fetch PPTX explicitly for version 1 again
    _, v1_reloaded = tmp_registry.get_template_version_pptx(res1.template.template_id, version_number=1)
    assert v1_reloaded.version_number == 1
    assert v1_reloaded.file_hash == v1_obj.file_hash

    # Fetch PPTX for active version (which is version 2)
    _, v2_active = tmp_registry.get_template_version_pptx(res1.template.template_id)
    assert v2_active.version_number == 2
    assert v2_active.file_hash != v1_obj.file_hash


