import uuid
from datetime import datetime
from typing import List, Optional
from pptx import Presentation

from app.schemas.template import TemplateInspectionResponse, SlideType
from app.schemas.fidelity import (
    TemplateFidelityManifest,
    SlideLayoutManifest,
    StaticRegionManifest,
    DynamicRegionManifest,
    RegionBBox,
)

EMU_PER_INCH = 914400.0


def generate_fidelity_manifest(
    inspection: TemplateInspectionResponse,
    file_hash: str,
    template_version: int,
    pptx_bytes: bytes,
) -> TemplateFidelityManifest:
    """Generate an immutable, version-specific fidelity manifest directly from the inspected PPTX."""
    prs = Presentation(io.BytesIO(pptx_bytes)) if isinstance(pptx_bytes, bytes) else pptx_bytes
    slide_w_in = prs.slide_width / EMU_PER_INCH
    slide_h_in = prs.slide_height / EMU_PER_INCH

    layouts: List[SlideLayoutManifest] = []

    for s_idx, slide_info in enumerate(inspection.slides):
        lm = slide_info.layout_model
        slide_type = lm.slide_type.value if lm else SlideType.CONTENT.value
        capacity = lm.capacity if lm else 1

        static_regions: List[StaticRegionManifest] = []
        dynamic_regions: List[DynamicRegionManifest] = []
        photo_regions: List[DynamicRegionManifest] = []

        # Map static text shapes
        if lm and lm.static_text_shapes:
            for st in lm.static_text_shapes:
                static_regions.append(
                    StaticRegionManifest(
                        shape_name=st.shape_name,
                        shape_id=st.shape_id,
                        shape_type=st.shape_type,
                        bbox=RegionBBox(
                            left_inches=st.left_inches,
                            top_inches=st.top_inches,
                            width_inches=st.width_inches,
                            height_inches=st.height_inches,
                        ),
                        has_text=True,
                        text_content=st.text,
                        is_image=False,
                    )
                )

        # Map static images
        if lm and lm.static_images:
            for si in lm.static_images:
                static_regions.append(
                    StaticRegionManifest(
                        shape_name=si.shape_name,
                        shape_id=si.shape_id,
                        shape_type="PICTURE",
                        bbox=RegionBBox(
                            left_inches=si.left_inches,
                            top_inches=si.top_inches,
                            width_inches=si.width_inches,
                            height_inches=si.height_inches,
                        ),
                        has_text=False,
                        is_image=True,
                    )
                )

        # Map dynamic slots
        if lm and lm.employee_slots:
            for slot in lm.employee_slots:
                # Text shapes in this slot
                for ts_idx, ts in enumerate(slot.text_shapes):
                    # Check if single multi-line card
                    lines = [ln.strip() for ln in ts.text.split("\n") if ln.strip()]
                    is_multiline = len(lines) >= 2
                    
                    if is_multiline:
                        # Card contains name, role, branch across lines
                        dynamic_regions.append(
                            DynamicRegionManifest(
                                region_id=f"slot_{slot.slot_index}_card",
                                slot_index=slot.slot_index,
                                field_key="employee_card",
                                shape_name=ts.shape_name,
                                shape_id=ts.shape_id,
                                bbox=RegionBBox(
                                    left_inches=ts.left_inches,
                                    top_inches=ts.top_inches,
                                    width_inches=ts.width_inches,
                                    height_inches=ts.height_inches,
                                ),
                                is_multiline_card=True,
                            )
                        )
                    else:
                        # Individual shape per field
                        dynamic_regions.append(
                            DynamicRegionManifest(
                                region_id=f"slot_{slot.slot_index}_field_{ts_idx}",
                                slot_index=slot.slot_index,
                                field_key="employee_name" if ts_idx == 0 else ("designation" if ts_idx == 1 else "branch"),
                                shape_name=ts.shape_name,
                                shape_id=ts.shape_id,
                                bbox=RegionBBox(
                                    left_inches=ts.left_inches,
                                    top_inches=ts.top_inches,
                                    width_inches=ts.width_inches,
                                    height_inches=ts.height_inches,
                                ),
                                is_multiline_card=False,
                            )
                        )

                # Image regions in this slot (photo frames)
                for p_idx, ir in enumerate(slot.image_regions):
                    p_reg = DynamicRegionManifest(
                        region_id=f"slot_{slot.slot_index}_photo_{p_idx}",
                        slot_index=slot.slot_index,
                        field_key="photo_placeholder",
                        shape_name=ir.shape_name,
                        shape_id=ir.shape_id,
                        bbox=RegionBBox(
                            left_inches=ir.left_inches,
                            top_inches=ir.top_inches,
                            width_inches=ir.width_inches,
                            height_inches=ir.height_inches,
                        ),
                    )
                    photo_regions.append(p_reg)
                    dynamic_regions.append(p_reg)

        layouts.append(
            SlideLayoutManifest(
                layout_variant_id=f"layout_slide_{s_idx + 1}",
                source_slide_index=s_idx,
                slide_type=slide_type,
                capacity=capacity,
                static_regions=static_regions,
                dynamic_regions=dynamic_regions,
                photo_regions=photo_regions,
            )
        )

    return TemplateFidelityManifest(
        manifest_id=uuid.uuid4().hex,
        template_id=inspection.template_id,
        template_version=template_version,
        source_file_hash=file_hash,
        slide_width_inches=round(slide_w_in, 3),
        slide_height_inches=round(slide_h_in, 3),
        aspect_ratio=inspection.aspect_ratio,
        total_source_slides=len(inspection.slides),
        layouts=layouts,
        created_at=datetime.utcnow().isoformat(),
    )
import io
