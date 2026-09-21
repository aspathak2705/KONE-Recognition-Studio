import openpyxl
from io import BytesIO
from typing import Dict, List, Tuple
from app.schemas.recognition import (
    RecognitionRecord,
    ValidationErrorItem,
    ExcelValidationResponse,
    FieldSchemaInfo,
    SchemaResponse,
)

HEADER_MAPPINGS: Dict[str, List[str]] = {
    "employee_name": ["employee_name", "employee name", "name", "emp_name", "full name", "full_name"],
    "designation": ["designation", "role", "title", "job_title", "job title", "position"],
    "branch": ["branch", "location", "office", "city", "unit", "branch_name"],
    "award_name": ["award_name", "award name", "award", "recognition", "category", "award_title"],
}

CANONICAL_FIELDS = list(HEADER_MAPPINGS.keys())


def get_schema_info() -> SchemaResponse:
    fields = [
        FieldSchemaInfo(
            field_name="employee_name",
            canonical_header="employee_name",
            accepted_aliases=HEADER_MAPPINGS["employee_name"],
            required=True,
            description="Employee's full legal or preferred name",
        ),
        FieldSchemaInfo(
            field_name="designation",
            canonical_header="designation",
            accepted_aliases=HEADER_MAPPINGS["designation"],
            required=True,
            description="Employee role or job title",
        ),
        FieldSchemaInfo(
            field_name="branch",
            canonical_header="branch",
            accepted_aliases=HEADER_MAPPINGS["branch"],
            required=True,
            description="Branch location or office name",
        ),
        FieldSchemaInfo(
            field_name="award_name",
            canonical_header="award_name",
            accepted_aliases=HEADER_MAPPINGS["award_name"],
            required=True,
            description="Quarterly award or recognition category name",
        ),
    ]
    return SchemaResponse(
        required_fields=fields,
        accepted_formats=[".xlsx"],
        rules=[
            "First row must contain header columns.",
            "Supported headers: employee_name, designation, branch, award_name.",
            "Values cannot be empty or whitespace-only.",
            "Duplicate records across all 4 fields will be flagged.",
        ],
    )


def normalize_header(header_text: str) -> str:
    cleaned = str(header_text).strip().lower().replace("-", "_").replace(" ", "_")
    for canonical, aliases in HEADER_MAPPINGS.items():
        normalized_aliases = [a.lower().replace("-", "_").replace(" ", "_") for a in aliases]
        if cleaned in normalized_aliases:
            return canonical
    return cleaned


def parse_and_validate_excel(file_content: bytes, filename: str) -> ExcelValidationResponse:
    try:
        wb = openpyxl.load_workbook(filename=BytesIO(file_content), data_only=True)
    except Exception as e:
        return ExcelValidationResponse(
            valid=False,
            filename=filename,
            total_rows=0,
            valid_rows=0,
            invalid_rows=0,
            duplicate_rows=0,
            errors=[ValidationErrorItem(row=0, column="file", message=f"Corrupted or invalid Excel file: {str(e)}")],
            records=[],
        )

    sheet = wb.active
    if sheet is None or sheet.max_row < 1:
        return ExcelValidationResponse(
            valid=False,
            filename=filename,
            total_rows=0,
            valid_rows=0,
            invalid_rows=0,
            duplicate_rows=0,
            errors=[ValidationErrorItem(row=0, column="sheet", message="Excel worksheet contains no data.")],
            records=[],
        )

    # Extract headers from row 1
    raw_headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    col_map: Dict[str, int] = {}  # canonical -> col_index (0-based)
    
    for idx, raw_val in enumerate(raw_headers):
        if raw_val is not None and str(raw_val).strip():
            normalized = normalize_header(str(raw_val))
            if normalized in CANONICAL_FIELDS and normalized not in col_map:
                col_map[normalized] = idx

    missing_fields = [f for f in CANONICAL_FIELDS if f not in col_map]
    if missing_fields:
        return ExcelValidationResponse(
            valid=False,
            filename=filename,
            total_rows=0,
            valid_rows=0,
            invalid_rows=0,
            duplicate_rows=0,
            errors=[
                ValidationErrorItem(
                    row=1,
                    column="header",
                    message=f"Missing required columns: {', '.join(missing_fields)}. Found headers: {[h for h in raw_headers if h]}",
                )
            ],
            records=[],
        )

    records: List[RecognitionRecord] = []
    errors: List[ValidationErrorItem] = []
    seen_combinations: set = set()
    duplicate_count = 0

    row_idx = 1
    for row_cells in sheet.iter_rows(min_row=2):
        row_idx += 1
        # Check if entire row is empty
        cell_values = [cell.value for cell in row_cells]
        if not any(v is not None and str(v).strip() for v in cell_values):
            continue  # Skip trailing blank lines

        row_data: Dict[str, str] = {}
        row_errors: List[str] = []

        for field in CANONICAL_FIELDS:
            c_idx = col_map[field]
            val = row_cells[c_idx].value if c_idx < len(row_cells) else None
            cleaned_val = str(val).strip() if val is not None else ""

            if not cleaned_val:
                msg = f"Required field '{field}' is missing or empty"
                row_errors.append(msg)
                errors.append(ValidationErrorItem(row=row_idx, column=field, message=msg))
                row_data[field] = ""
            else:
                row_data[field] = cleaned_val

        # Duplicate check across non-empty complete records
        if not row_errors:
            combo_key = (
                row_data["employee_name"].lower(),
                row_data["designation"].lower(),
                row_data["branch"].lower(),
                row_data["award_name"].lower(),
            )
            if combo_key in seen_combinations:
                msg = f"Duplicate record detected for '{row_data['employee_name']}' with award '{row_data['award_name']}'"
                row_errors.append(msg)
                errors.append(ValidationErrorItem(row=row_idx, column="duplicate", message=msg))
                duplicate_count += 1
            else:
                seen_combinations.add(combo_key)

        record = RecognitionRecord(
            row_number=row_idx,
            employee_name=row_data.get("employee_name", ""),
            designation=row_data.get("designation", ""),
            branch=row_data.get("branch", ""),
            award_name=row_data.get("award_name", ""),
            is_valid=len(row_errors) == 0,
            errors=row_errors,
        )
        records.append(record)

    total_rows = len(records)
    valid_rows = sum(1 for r in records if r.is_valid)
    invalid_rows = total_rows - valid_rows

    return ExcelValidationResponse(
        valid=(invalid_rows == 0 and total_rows > 0),
        filename=filename,
        total_rows=total_rows,
        valid_rows=valid_rows,
        invalid_rows=invalid_rows,
        duplicate_rows=duplicate_count,
        errors=errors,
        records=records,
    )
