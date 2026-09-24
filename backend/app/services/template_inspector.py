import uuid
from io import BytesIO
from typing import List
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.schemas.template import (
    TextShapeInfo,
    SlideInfo,
    TemplateInspectionResponse,
)

# 1 inch = 914400 EMUs
EMU_PER_INCH = 914400.0


def inspect_powerpoint_template(file_content: bytes, filename: str) -> TemplateInspectionResponse:
    template_id = uuid.uuid4().hex
    warnings: List[str] = []

    try:
        prs = Presentation(BytesIO(file_content))
    except Exception as e:
        return TemplateInspectionResponse(
            valid=False,
            template_id=template_id,
            filename=filename,
            slide_count=0,
            slide_width_inches=0.0,
            slide_height_inches=0.0,
            aspect_ratio="unknown",
            slides=[],
            warnings=[f"Corrupted or invalid PowerPoint presentation: {str(e)}"],
        )

    width_in = round(prs.slide_width / EMU_PER_INCH, 2)
    height_in = round(prs.slide_height / EMU_PER_INCH, 2)
    
    # Calculate aspect ratio
    ratio_val = width_in / height_in if height_in > 0 else 0
    if abs(ratio_val - (16 / 9)) < 0.1:
        aspect_ratio = "16:9"
    elif abs(ratio_val - (4 / 3)) < 0.1:
        aspect_ratio = "4:3"
    else:
        aspect_ratio = f"{width_in}:{height_in}"

    slides_info: List[SlideInfo] = []
    total_placeholders = 0
    total_shapes = 0
    total_text_shapes = 0
    total_tables = 0
    total_groups = 0

    def process_shape_list(shapes, text_shapes_list, stats_dict):
        for shape in shapes:
            stats_dict["total_shapes"] += 1

            # Check if group shape
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                stats_dict["total_groups"] += 1
                try:
                    process_shape_list(shape.shapes, text_shapes_list, stats_dict)
                except Exception:
                    pass
                continue

            # Check if table
            if shape.has_table:
                stats_dict["total_tables"] += 1
                continue

            is_text = shape.has_text_frame
            text_content = shape.text.strip() if is_text and shape.text else ""
            
            is_ph = shape.is_placeholder
            ph_type_str = None
            if is_ph:
                try:
                    ph_type_str = str(shape.placeholder_format.type)
                except Exception:
                    ph_type_str = "PLACEHOLDER"

            if is_text or is_ph:
                stats_dict["total_text_shapes"] += 1
                shape_name = shape.name or f"Shape_{shape.shape_id}"
                shape_type_str = str(shape.shape_type) if hasattr(shape, "shape_type") else "TEXT"

                text_shapes_list.append(
                    TextShapeInfo(
                        shape_name=shape_name,
                        shape_type=shape_type_str,
                        text=text_content,
                        left_inches=round(shape.left / EMU_PER_INCH, 2),
                        top_inches=round(shape.top / EMU_PER_INCH, 2),
                        width_inches=round(shape.width / EMU_PER_INCH, 2),
                        height_inches=round(shape.height / EMU_PER_INCH, 2),
                        unit="inches",
                        is_placeholder=is_ph,
                        placeholder_type=ph_type_str,
                    )
                )

    for idx, slide in enumerate(prs.slides, start=1):
        text_shapes: List[TextShapeInfo] = []
        shape_count = len(slide.shapes)
        placeholder_count = len(slide.placeholders)
        total_placeholders += placeholder_count

        slide_stats = {
            "total_shapes": 0,
            "total_text_shapes": 0,
            "total_tables": 0,
            "total_groups": 0,
        }

        process_shape_list(slide.shapes, text_shapes, slide_stats)

        total_shapes += slide_stats["total_shapes"]
        total_text_shapes += slide_stats["total_text_shapes"]
        total_tables += slide_stats["total_tables"]
        total_groups += slide_stats["total_groups"]

        slides_info.append(
            SlideInfo(
                slide_number=idx,
                shape_count=slide_stats["total_shapes"],
                text_shapes=text_shapes,
                placeholders_count=placeholder_count,
                tables_count=slide_stats["total_tables"],
                groups_count=slide_stats["total_groups"],
            )
        )

    if total_placeholders == 0:
        warnings.append(
            "No standard PowerPoint placeholders found. Template elements rely on text shapes or custom frames."
        )

    if total_groups > 0:
        warnings.append(
            f"Found {total_groups} grouped shape(s). Grouped shapes were recursively inspected for text frames."
        )

    warnings.append(
        "No dynamic field mapping established yet. Manual mapping is required before Phase 2 generation."
    )

    if len(slides_info) == 0:
        warnings.append("Presentation contains no slides.")

    return TemplateInspectionResponse(
        valid=True,
        template_id=template_id,
        filename=filename,
        slide_count=len(slides_info),
        slide_width_inches=width_in,
        slide_height_inches=height_in,
        unit="inches",
        aspect_ratio=aspect_ratio,
        total_shapes_count=total_shapes,
        total_text_shapes_count=total_text_shapes,
        total_placeholders_count=total_placeholders,
        total_tables_count=total_tables,
        total_groups_count=total_groups,
        slides=slides_info,
        warnings=warnings,
    )

