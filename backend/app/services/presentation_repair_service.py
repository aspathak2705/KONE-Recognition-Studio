import io
import copy
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple, Dict, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.schemas.fidelity import (
    TemplateFidelityManifest,
    ValidationFailureDetail,
    RepairRecord,
    RepairOutcome,
    QualityClassification,
    DeliveryStatus,
    FailureSeverity,
)

EMU_PER_INCH = 914400.0


class PresentationRepairService:
    """
    Template-independent, bounded automatic repair pipeline for generated PowerPoint presentations.
    Repairs genuinely repairable structural and layout defects once using evidence from the
    immutable master template and fidelity manifest.
    """

    SUPPORTED_REPAIR_TYPES = {
        "MISSING_PHOTO_REGION",
        "MISSING_STATIC_ASSET",
        "GEOMETRY_DEVIATION",
    }

    @classmethod
    def can_attempt_repair(
        cls,
        failures: List[ValidationFailureDetail],
        current_attempt_count: int,
    ) -> bool:
        """
        Check if repair is permissible and supported:
        1. Strict attempt limit: exactly 0 prior attempts (at most 1 repair per generation).
        2. At least one supported, repairable failure exists.
        3. No irrecoverable/unsupported critical failures.
        """
        if current_attempt_count >= 1:
            return False

        if not failures:
            return False

        # Any failure with is_repairable=False and CRITICAL severity cannot be safely repaired
        for f in failures:
            if not f.is_repairable and f.severity == FailureSeverity.CRITICAL:
                return False

        # Must have at least one supported repair type
        repairable_supported = [
            f for f in failures
            if f.is_repairable and f.failure_type in cls.SUPPORTED_REPAIR_TYPES
        ]
        return len(repairable_supported) > 0

    @classmethod
    def repair_presentation(
        cls,
        generated_pptx_bytes: bytes,
        source_template_bytes: bytes,
        manifest: TemplateFidelityManifest,
        failures: List[ValidationFailureDetail],
    ) -> Tuple[bytes, RepairRecord]:
        """
        Execute bounded repair operation on the generated PPTX binary.
        Returns: (repaired_pptx_bytes, repair_record)
        """
        now_str = datetime.utcnow().isoformat()
        repair_record = RepairRecord(
            attempt_count=1,
            repaired_timestamp=now_str,
            post_repair_verified=False,
        )

        try:
            gen_prs = Presentation(io.BytesIO(generated_pptx_bytes))
            src_prs = Presentation(io.BytesIO(source_template_bytes))
        except Exception as e:
            repair_record.outcome = RepairOutcome.FAILED
            repair_record.details = f"Cannot parse presentations for repair: {str(e)}"
            return generated_pptx_bytes, repair_record

        addressed_types = set()
        repair_actions: List[str] = []

        # Group failures by slide_number (1-indexed)
        failures_by_slide: Dict[int, List[ValidationFailureDetail]] = {}
        for f in failures:
            failures_by_slide.setdefault(f.slide_number, []).append(f)

        # Iterate over generated slides
        for slide_num, slide_failures in failures_by_slide.items():
            slide_idx = slide_num - 1
            if slide_idx < 0 or slide_idx >= len(gen_prs.slides):
                continue

            gen_slide = gen_prs.slides[slide_idx]

            # Match corresponding layout in manifest
            matched_layout = None
            if slide_idx < len(manifest.layouts):
                matched_layout = manifest.layouts[slide_idx]
            elif manifest.layouts:
                matched_layout = manifest.layouts[0]

            if not matched_layout:
                continue

            # Determine source slide in master template
            src_slide_idx = matched_layout.source_slide_index
            if src_slide_idx >= len(src_prs.slides):
                src_slide_idx = 0
            src_slide = src_prs.slides[src_slide_idx]

            for fail in slide_failures:
                if fail.failure_type == "MISSING_PHOTO_REGION":
                    success, detail = cls._repair_missing_photo_region(
                        gen_slide=gen_slide,
                        src_slide=src_slide,
                        target_region_name=fail.region,
                        matched_layout=matched_layout,
                    )
                    if success:
                        addressed_types.add("MISSING_PHOTO_REGION")
                        repair_actions.append(f"Slide #{slide_num}: {detail}")

                elif fail.failure_type == "MISSING_STATIC_ASSET":
                    success, detail = cls._repair_missing_static_asset(
                        gen_slide=gen_slide,
                        src_slide=src_slide,
                        target_region_name=fail.region,
                    )
                    if success:
                        addressed_types.add("MISSING_STATIC_ASSET")
                        repair_actions.append(f"Slide #{slide_num}: {detail}")

                elif fail.failure_type == "GEOMETRY_DEVIATION":
                    success, detail = cls._repair_geometry_deviation(
                        gen_slide=gen_slide,
                        target_region_name=fail.region,
                        matched_layout=matched_layout,
                    )
                    if success:
                        addressed_types.add("GEOMETRY_DEVIATION")
                        repair_actions.append(f"Slide #{slide_num}: {detail}")

        if addressed_types:
            # Save repaired presentation
            out_bio = io.BytesIO()
            gen_prs.save(out_bio)
            repaired_bytes = out_bio.getvalue()

            repair_record.outcome = RepairOutcome.SUCCESS
            repair_record.strategy_used = "+".join(sorted(addressed_types))
            repair_record.finding_types_addressed = sorted(list(addressed_types))
            repair_record.details = "; ".join(repair_actions)
            return repaired_bytes, repair_record
        else:
            repair_record.outcome = RepairOutcome.UNSUPPORTED
            repair_record.details = "No repairable failures could be deterministically matched to source template regions."
            return generated_pptx_bytes, repair_record

    @classmethod
    def _find_shape_in_slide(cls, shapes, name: str):
        for s in shapes:
            if s.name == name:
                return s
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                nested = cls._find_shape_in_slide(s.shapes, name)
                if nested:
                    return nested
        return None

    @classmethod
    def _repair_missing_photo_region(
        cls,
        gen_slide,
        src_slide,
        target_region_name: str,
        matched_layout,
    ) -> Tuple[bool, str]:
        """
        Restore missing photo frame/region from source template slide.
        Deep-copies the picture/shape from the master template, binds relationships,
        and clears sample bitmap so HR can insert photo manually while preserving frame.
        """
        # 1. Locate photo region definition in manifest
        photo_manifest = next((pr for pr in matched_layout.photo_regions if pr.shape_name == target_region_name), None)
        if not photo_manifest:
            return False, f"Photo region '{target_region_name}' not found in layout manifest."

        # 2. Check if a shape already exists at this exact bbox (avoid duplicate insertion)
        for s in gen_slide.shapes:
            s_l = s.left / EMU_PER_INCH
            s_t = s.top / EMU_PER_INCH
            if abs(s_l - photo_manifest.bbox.left_inches) <= 0.15 and abs(s_t - photo_manifest.bbox.top_inches) <= 0.15:
                # Shape is actually present at coordinates
                return True, f"Photo region '{target_region_name}' verified present at coordinates."

        # 3. Locate source shape on master template slide
        src_shape = cls._find_shape_in_slide(src_slide.shapes, target_region_name)
        if not src_shape:
            # Fallback: find any shape at the matching bounding box in the source slide
            for s in src_slide.shapes:
                s_l = s.left / EMU_PER_INCH
                s_t = s.top / EMU_PER_INCH
                if abs(s_l - photo_manifest.bbox.left_inches) <= 0.2 and abs(s_t - photo_manifest.bbox.top_inches) <= 0.2:
                    src_shape = s
                    break

        if not src_shape:
            return False, f"Source shape for photo region '{target_region_name}' not found in master template."

        # 4. Clone source shape into generated slide with relationships
        elem = copy.deepcopy(src_shape._element)
        if src_shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blips = elem.xpath('.//a:blip')
            if blips:
                old_rid = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if old_rid and old_rid in src_slide.part.rels:
                    rel = src_slide.part.rels[old_rid]
                    new_rid = gen_slide.part.relate_to(rel.target_part, rel.reltype)
                    blips[0].set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', new_rid)

        gen_slide.shapes._spTree.append(elem)

        # 5. Clear sample image content to neutral blank frame while preserving geometry
        cloned_shape = cls._find_shape_in_slide(gen_slide.shapes, src_shape.name)
        if cloned_shape and cloned_shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blips = cloned_shape._element.xpath('.//a:blip')
            if blips:
                rId = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if rId and rId in gen_slide.part.rels:
                    try:
                        from PIL import Image
                        part = gen_slide.part.related_part(rId)
                        blank = Image.new('RGB', (100, 100), (248, 248, 248))
                        bio = io.BytesIO()
                        blank.save(bio, format='PNG')
                        part._blob = bio.getvalue()
                    except Exception:
                        pass

        return True, f"Restored photo region frame '{target_region_name}' from master template."

    @classmethod
    def _repair_missing_static_asset(
        cls,
        gen_slide,
        src_slide,
        target_region_name: str,
    ) -> Tuple[bool, str]:
        """Restore missing static graphic or banner shape from source template slide."""
        src_shape = cls._find_shape_in_slide(src_slide.shapes, target_region_name)
        if not src_shape:
            return False, f"Static asset '{target_region_name}' not found on source template slide."

        elem = copy.deepcopy(src_shape._element)
        if src_shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            blips = elem.xpath('.//a:blip')
            if blips:
                old_rid = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if old_rid and old_rid in src_slide.part.rels:
                    rel = src_slide.part.rels[old_rid]
                    new_rid = gen_slide.part.relate_to(rel.target_part, rel.reltype)
                    blips[0].set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', new_rid)

        gen_slide.shapes._spTree.append(elem)
        return True, f"Restored static asset '{target_region_name}' from master template."

    @classmethod
    def _repair_geometry_deviation(
        cls,
        gen_slide,
        target_region_name: str,
        matched_layout,
    ) -> Tuple[bool, str]:
        """Align shape back to manifest-defined coordinates if displaced."""
        target_shape = cls._find_shape_in_slide(gen_slide.shapes, target_region_name)
        if not target_shape:
            return False, f"Shape '{target_region_name}' not found on slide for geometry adjustment."

        # Search for region bounding box in manifest
        reg = next((r for r in matched_layout.photo_regions if r.shape_name == target_region_name), None)
        if not reg:
            reg = next((r for r in matched_layout.static_regions if r.shape_name == target_region_name), None)

        if not reg:
            return False, f"Region definition for '{target_region_name}' not found in manifest."

        target_shape.left = int(reg.bbox.left_inches * EMU_PER_INCH)
        target_shape.top = int(reg.bbox.top_inches * EMU_PER_INCH)
        target_shape.width = int(reg.bbox.width_inches * EMU_PER_INCH)
        target_shape.height = int(reg.bbox.height_inches * EMU_PER_INCH)
        return True, f"Restored geometry for '{target_region_name}' to manifest coordinates."
